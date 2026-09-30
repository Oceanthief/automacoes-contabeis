import shutil
import sys
from pathlib import Path

import openpyxl
import pytest

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))

import main  # noqa: E402

EXTRATO_EXEMPLO = RAIZ / "dados_exemplo" / "extrato_2026_08.xls"


@pytest.fixture
def isolar_diretorio_saida(tmp_path, monkeypatch):
    """Evita que os testes escrevam dentro de output/ do projeto de verdade."""
    monkeypatch.setattr(main, "DIRETORIO_SAIDA_BASE", tmp_path / "output")
    return tmp_path / "output"


def test_execucao_completa_com_arquivo_de_exemplo(isolar_diretorio_saida, tmp_path):
    copia = tmp_path / "extrato_do_mes.xls"
    shutil.copy(EXTRATO_EXEMPLO, copia)

    codigo = main.main([str(copia)])
    assert codigo == 0

    caminho_xlsx = isolar_diretorio_saida / "2026-08" / "contamax_2026-08.xlsx"
    caminho_txt = isolar_diretorio_saida / "2026-08" / "contamax_2026-08.txt"
    assert caminho_xlsx.exists()
    assert caminho_txt.exists()

    wb = openpyxl.load_workbook(str(caminho_xlsx))
    assert wb.sheetnames == ["LANCAMENTOS", "AUDITORIA_DIARIA"]
    assert wb["LANCAMENTOS"].max_row - 1 == 11  # 12 dias com movimento, 1 empate sem lançamento

    conteudo_txt = caminho_txt.read_bytes()
    assert conteudo_txt.count(b"\r\n") == 11


def test_nao_sobrescreve_resultado_existente(isolar_diretorio_saida, tmp_path):
    copia = tmp_path / "extrato.xls"
    shutil.copy(EXTRATO_EXEMPLO, copia)

    assert main.main([str(copia)]) == 0
    codigo_segunda_vez = main.main([str(copia)])
    assert codigo_segunda_vez == 1  # deve recusar, não sobrescrever


def test_arquivo_inexistente_da_erro_amigavel(isolar_diretorio_saida, capsys):
    codigo = main.main(["/caminho/que/nao/existe.xls"])
    assert codigo == 1
    saida = capsys.readouterr().out
    assert "Não encontrei o arquivo" in saida
    assert "Traceback" not in saida  # nada de stack trace pra usuário não técnico


def test_extensao_nao_suportada_da_erro_amigavel(isolar_diretorio_saida, tmp_path, capsys):
    arquivo_xlsx = tmp_path / "extrato.xlsx"
    arquivo_xlsx.write_text("qualquer coisa")

    codigo = main.main([str(arquivo_xlsx)])
    assert codigo == 1
    saida = capsys.readouterr().out
    assert ".xls" in saida


def test_sem_argumentos_mostra_uso(capsys):
    codigo = main.main([])
    assert codigo == 2
    saida = capsys.readouterr().out
    assert "Uso:" in saida
