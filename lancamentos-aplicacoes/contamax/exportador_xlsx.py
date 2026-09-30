"""
Geração do Excel mensal de conferência (contamax_AAAA-MM.xlsx).

Este módulo só produz um arquivo de CONFERÊNCIA — números prontos, sem
fórmula nenhuma (conforme pedido explicitamente: "nenhuma fórmula externa").
Isso é diferente de um modelo financeiro editável: aqui o objetivo é alguém
bater o olho e confirmar que o cálculo bate antes de importar no ERP, não
editar a planilha e ela recalcular. Por isso os valores são gravados como
números finais, não como `=SOMA(...)`.
"""
from __future__ import annotations

from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import LancamentoContamax, StatusLancamento
from .regras import HISTORICO_APLICACAO, HISTORICO_RESGATE
from .resumo import montar_resumo
from .validacoes import validar_resultados

FORMATO_DATA = "DD/MM/YYYY"
FORMATO_MOEDA = "R$ #,##0.00"

FONTE_CABECALHO = Font(name="Arial", bold=True, color="FFFFFF")
PREENCHIMENTO_CABECALHO = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
FONTE_CORPO = Font(name="Arial")


def nome_arquivo_padrao(resultados: list[LancamentoContamax]) -> str:
    """Deriva 'contamax_AAAA-MM.xlsx' do período dos resultados."""
    resumo = montar_resumo(resultados)
    return f"contamax_{resumo.periodo}.xlsx"


def _regra_aplicada_texto(r: LancamentoContamax) -> str:
    if r.status == StatusLancamento.GERADO and r.historico == HISTORICO_RESGATE:
        return "DÉBITO > CRÉDITO — RESGATE"
    if r.status == StatusLancamento.GERADO and r.historico == HISTORICO_APLICACAO:
        return "CRÉDITO > DÉBITO — APLICAÇÃO"
    if r.status == StatusLancamento.SEM_LANCAMENTO_EMPATE:
        return "DÉBITO = CRÉDITO — SEM LANÇAMENTO"
    if r.status == StatusLancamento.ERRO_VALOR_INVALIDO:
        return "VALOR INVÁLIDO NO DIA — PENDENTE DE REVISÃO"
    # Não deveria ser alcançável com os status atuais, mas evita gravar uma
    # linha muda no Excel se um status novo for adicionado no futuro sem
    # atualizar esta função.
    return f"STATUS NÃO MAPEADO: {r.status.value}"


def _escrever_cabecalho(ws: Worksheet, colunas: list[str]) -> None:
    for col_idx, titulo in enumerate(colunas, start=1):
        cel = ws.cell(row=1, column=col_idx, value=titulo)
        cel.font = FONTE_CABECALHO
        cel.fill = PREENCHIMENTO_CABECALHO
        cel.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"


def _ajustar_largura_colunas(ws: Worksheet, larguras: list[int]) -> None:
    for idx, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = largura


def _preencher_aba_lancamentos(ws: Worksheet, resultados: list[LancamentoContamax]) -> None:
    _escrever_cabecalho(ws, ["DATA", "CONTA DEBITO", "CONTA CREDITO", "HISTORICO", "VALOR"])
    gerados = sorted(
        (r for r in resultados if r.status == StatusLancamento.GERADO), key=lambda r: r.data
    )
    for linha_idx, r in enumerate(gerados, start=2):
        ws.cell(row=linha_idx, column=1, value=r.data).number_format = FORMATO_DATA
        ws.cell(row=linha_idx, column=2, value=r.conta_debito)
        ws.cell(row=linha_idx, column=3, value=r.conta_credito)
        ws.cell(row=linha_idx, column=4, value=r.historico)
        ws.cell(row=linha_idx, column=5, value=float(r.valor)).number_format = FORMATO_MOEDA
        for col in range(1, 6):
            ws.cell(row=linha_idx, column=col).font = FONTE_CORPO
    _ajustar_largura_colunas(ws, [12, 14, 14, 32, 14])


def _preencher_aba_auditoria(ws: Worksheet, resultados: list[LancamentoContamax]) -> None:
    _escrever_cabecalho(
        ws, ["DATA", "TOTAL DEBITO", "TOTAL CREDITO", "DIFERENCA", "REGRA APLICADA", "STATUS"]
    )
    ordenados = sorted(resultados, key=lambda r: r.data)
    for linha_idx, r in enumerate(ordenados, start=2):
        ws.cell(row=linha_idx, column=1, value=r.data).number_format = FORMATO_DATA
        ws.cell(row=linha_idx, column=2, value=float(r.total_debito)).number_format = FORMATO_MOEDA
        ws.cell(row=linha_idx, column=3, value=float(r.total_credito)).number_format = FORMATO_MOEDA
        ws.cell(row=linha_idx, column=4, value=float(r.diferenca)).number_format = FORMATO_MOEDA
        ws.cell(row=linha_idx, column=5, value=_regra_aplicada_texto(r))
        ws.cell(row=linha_idx, column=6, value=r.status.value)
        for col in range(1, 7):
            ws.cell(row=linha_idx, column=col).font = FONTE_CORPO
    _ajustar_largura_colunas(ws, [12, 15, 15, 14, 32, 24])


def gerar_excel_conferencia(resultados: list[LancamentoContamax], caminho_saida: str) -> None:
    """
    Gera o arquivo `caminho_saida` com as abas LANCAMENTOS e
    AUDITORIA_DIARIA. Levanta ErroCriticoValidacao (contamax.validacoes) e
    NÃO grava nada em disco se qualquer checagem de sanidade falhar — a
    validação roda antes de qualquer escrita de arquivo.
    """
    validar_resultados(resultados)

    wb = Workbook()
    ws_lancamentos = wb.active
    ws_lancamentos.title = "LANCAMENTOS"
    _preencher_aba_lancamentos(ws_lancamentos, resultados)

    ws_auditoria = wb.create_sheet("AUDITORIA_DIARIA")
    _preencher_aba_auditoria(ws_auditoria, resultados)

    wb.save(caminho_saida)
