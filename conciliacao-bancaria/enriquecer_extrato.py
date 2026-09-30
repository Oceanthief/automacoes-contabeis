"""Enriquece o extrato padronizado (saída da Fase 1) com a coluna
NOVA DESCRICAO e a classificação de tarifas bancárias.

Escopo: só isso. SISPAG FORNECEDORES fica com NOVA DESCRICAO vazia — será
preenchida na próxima fase a partir de auditoria_sispag.xlsx. Não mexe em
DESCRICAO, DATA, VALOR ou SEQUENCIA; não faz DE-PARA de fornecedor; não gera
TXT; não sobrescreve o arquivo de entrada.

Uso:
    python enriquecer_extrato.py <saida.xlsx> <saida_enriquecida.xlsx>
"""
from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from enriquecimento.classificacao_extrato import (
    CONTA_TARIFA,
    DESCRICAO_SISPAG_FORNECEDORES,
    ClassificacaoError,
    classificar_linha,
)
from leitura.extrato_padronizado import ExtratoPadronizadoError, ler_extrato_padronizado

COLUNAS_SAIDA = [
    "col_sem_info_1",
    "DEBITO",
    "CREDITO",
    "DATA",
    "VALOR",
    "col_sem_info_2",
    "DESCRICAO",
    "NOVA DESCRICAO",
    "SEQUENCIA",
]


class EnriquecimentoError(Exception):
    """Erros de validação específicos desta fase."""


def enriquecer(df) -> list[dict]:
    """Aplica classificar_linha a cada movimentação, preservando a ordem original."""
    linhas_saida = []
    for _, linha in df.iterrows():
        try:
            resultado = classificar_linha(linha["DESCRICAO"], linha["DEBITO"], linha["CREDITO"])
        except ClassificacaoError as exc:
            raise EnriquecimentoError(
                f"SEQUENCIA {linha['SEQUENCIA']}: {exc}"
            ) from exc

        linhas_saida.append(
            {
                "col_sem_info_1": "",
                "DEBITO": resultado.debito,
                "CREDITO": resultado.credito,
                "DATA": linha["DATA"],
                "VALOR": linha["VALOR"],
                "col_sem_info_2": "",
                "DESCRICAO": linha["DESCRICAO"],
                "NOVA DESCRICAO": resultado.nova_descricao,
                "SEQUENCIA": int(linha["SEQUENCIA"]),
                "_eh_pagamento": resultado.eh_pagamento,
                "_eh_tarifa": resultado.eh_tarifa,
            }
        )
    return linhas_saida


def validar(df_original, linhas_saida: list[dict]) -> None:
    """Critérios de aceite 1, 2, 3 e 9 — falha alto e claro, nunca silenciosa."""
    if len(linhas_saida) != len(df_original):
        raise EnriquecimentoError(
            f"Quantidade de linhas mudou: entrada={len(df_original)}, saída={len(linhas_saida)}"
        )

    sequencias_entrada = set(df_original["SEQUENCIA"].astype(int))
    sequencias_saida = {l["SEQUENCIA"] for l in linhas_saida}
    if sequencias_entrada != sequencias_saida:
        raise EnriquecimentoError("SEQUENCIA da saída não bate com a da entrada.")

    descricoes_entrada = list(df_original["DESCRICAO"])
    descricoes_saida = [l["DESCRICAO"] for l in linhas_saida]
    if descricoes_entrada != descricoes_saida:
        raise EnriquecimentoError("DESCRICAO original foi alterada em alguma linha.")

    for l in linhas_saida:
        if l["_eh_tarifa"]:
            if l["DEBITO"] != CONTA_TARIFA:
                raise EnriquecimentoError(f"SEQUENCIA {l['SEQUENCIA']}: tarifa sem DEBITO={CONTA_TARIFA}.")
            if l["CREDITO"] != "10001":
                raise EnriquecimentoError(f"SEQUENCIA {l['SEQUENCIA']}: tarifa sem CREDITO=10001.")


def gravar_excel(linhas_saida: list[dict], caminho_saida: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Extrato Enriquecido"

    fonte_cabecalho = Font(name="Arial", size=10, bold=True)
    fonte_dado = Font(name="Arial", size=10)

    rotulos = ["SEM INFORMACAO", "DEBITO", "CREDITO", "DATA", "VALOR", "SEM INFORMACAO", "DESCRICAO", "NOVA DESCRICAO", "SEQUENCIA"]
    ws.append(rotulos)
    for celula in ws[1]:
        celula.font = fonte_cabecalho

    for l in linhas_saida:
        ws.append(
            [
                l["col_sem_info_1"],
                l["DEBITO"],
                l["CREDITO"],
                l["DATA"],
                l["VALOR"],
                l["col_sem_info_2"],
                l["DESCRICAO"],
                l["NOVA DESCRICAO"],
                l["SEQUENCIA"],
            ]
        )

    for linha_celulas in ws.iter_rows(min_row=2):
        for celula in linha_celulas:
            celula.font = fonte_dado
        # DEBITO e CREDITO são códigos de conta, não números: forçados a
        # texto para "00248" nunca virar 248 ao reabrir no Excel.
        linha_celulas[1].number_format = "@"
        linha_celulas[1].value = str(linha_celulas[1].value) if linha_celulas[1].value != "" else ""
        linha_celulas[2].number_format = "@"
        linha_celulas[2].value = str(linha_celulas[2].value) if linha_celulas[2].value != "" else ""
        linha_celulas[3].number_format = "DD/MM/YYYY"  # DATA
        linha_celulas[4].number_format = "#,##0.00"  # VALOR

    for coluna_celulas in ws.columns:
        largura = max((len(str(c.value)) for c in coluna_celulas if c.value is not None), default=0)
        ws.column_dimensions[coluna_celulas[0].column_letter].width = min(largura + 2, 60)

    caminho_saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho_saida)


def _imprimir_resumo(df_original, linhas_saida: list[dict]) -> None:
    qtd_pagamentos = sum(1 for l in linhas_saida if l["_eh_pagamento"])
    qtd_recebimentos = sum(1 for l in linhas_saida if not l["_eh_pagamento"])
    qtd_tarifas = sum(1 for l in linhas_saida if l["_eh_tarifa"])
    qtd_sispag = sum(1 for l in linhas_saida if l["DESCRICAO"] == DESCRICAO_SISPAG_FORNECEDORES)

    descricoes_tarifa = sorted({l["DESCRICAO"] for l in linhas_saida if l["_eh_tarifa"]})

    print(f"Quantidade de pagamentos:              {qtd_pagamentos}")
    print(f"Quantidade de recebimentos:            {qtd_recebimentos}")
    print(f"Quantidade de tarifas encontradas:     {qtd_tarifas}")
    print(f"Quantidade de SISPAG (próxima fase):   {qtd_sispag}")
    print()
    print("Descrições classificadas como tarifa (distintas):")
    for d in descricoes_tarifa:
        print(f"  - {d}")


def executar(caminho_entrada: str, caminho_saida: str) -> list[dict]:
    entrada = Path(caminho_entrada)
    saida = Path(caminho_saida)

    if saida.resolve() == entrada.resolve():
        raise EnriquecimentoError("O arquivo de saída não pode ser igual ao arquivo de entrada.")

    df = ler_extrato_padronizado(entrada)
    linhas_saida = enriquecer(df)
    validar(df, linhas_saida)
    gravar_excel(linhas_saida, saida)

    _imprimir_resumo(df, linhas_saida)
    print(f"\nOK: {len(linhas_saida)} movimentações gravadas em {saida}")

    return linhas_saida


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Uso: python enriquecer_extrato.py <saida.xlsx> <saida_enriquecida.xlsx>")
        sys.exit(1)
    try:
        executar(sys.argv[1], sys.argv[2])
    except (ExtratoPadronizadoError, EnriquecimentoError) as exc:
        print(f"ERRO: {exc}")
        sys.exit(1)
