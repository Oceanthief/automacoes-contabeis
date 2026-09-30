"""Concilia os lançamentos SISPAG FORNECEDORES do extrato padronizado com os
pagamentos únicos reconstruídos de Contas Pagas, e gera um Excel de auditoria.

Escopo: só identificação e auditoria. Não altera a DESCRICAO do extrato, não
gera TXT de importação, não preenche conta contábil — isso fica para a
próxima fase, depois de revisão humana dos casos AMBIGUO/NAO_ENCONTRADO.

Uso:
    python conciliar_sispag.py <extrato_padronizado.xlsx> <contas_pagas.xlsx> <auditoria_sispag.xlsx>
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from conciliacao.motor_sispag import (
    STATUS_AMBIGUO,
    STATUS_EXATO_1X1,
    STATUS_EXATO_SOMA,
    STATUS_NAO_ENCONTRADO,
    ResultadoSispag,
    conciliar,
)
from leitura.contas_pagas import ContasPagasError, ler_contas_pagas, montar_pagamentos
from leitura.extrato_padronizado import (
    ExtratoPadronizadoError,
    ler_extrato_padronizado,
    listar_sispag_fornecedores,
)
from normalizacao import valor_para_decimal

COLUNAS_AUDITORIA = [
    "SEQUENCIA",
    "DATA",
    "VALOR_SISPAG",
    "STATUS",
    "TOTAL_ENCONTRADO",
    "DIFERENCA",
    "QTD_PAGAMENTOS",
    "TITULOS",
    "NOMES",
    "NF_NUM",
    "VALORES",
    "FORMA_PGTO",
    "HISTORICOS",
]


def _formatar_lista(pagamentos, campo: str) -> str:
    extrator = {
        "titulo": lambda p: str(p.titulo),
        "nome": lambda p: p.nome,
        "nf_num": lambda p: p.nf_num or "",
        "valor": lambda p: f"{p.valor_classificado:.2f}",
        "forma_pgto": lambda p: p.forma_pgto,
        "historico": lambda p: p.historico,
    }[campo]
    return "; ".join(extrator(p) for p in pagamentos)


def _formatar_alternativas(alternativas: list[list], campo: str) -> str:
    return " OU ".join(_formatar_lista(combo, campo) for combo in alternativas)


def _linha_auditoria(r: ResultadoSispag) -> dict:
    if r.status in (STATUS_EXATO_1X1, STATUS_EXATO_SOMA):
        titulos = _formatar_lista(r.pagamentos, "titulo")
        nomes = _formatar_lista(r.pagamentos, "nome")
        nf_num = _formatar_lista(r.pagamentos, "nf_num")
        valores = _formatar_lista(r.pagamentos, "valor")
        forma_pgto = _formatar_lista(r.pagamentos, "forma_pgto")
        historicos = _formatar_lista(r.pagamentos, "historico")
        qtd = len(r.pagamentos)
    elif r.status == STATUS_AMBIGUO:
        titulos = _formatar_alternativas(r.alternativas, "titulo")
        nomes = _formatar_alternativas(r.alternativas, "nome")
        nf_num = _formatar_alternativas(r.alternativas, "nf_num")
        valores = _formatar_alternativas(r.alternativas, "valor")
        forma_pgto = _formatar_alternativas(r.alternativas, "forma_pgto")
        historicos = _formatar_alternativas(r.alternativas, "historico")
        qtd = 0
    else:  # NAO_ENCONTRADO
        titulos = nomes = nf_num = valores = forma_pgto = historicos = ""
        qtd = 0

    return {
        "SEQUENCIA": r.sequencia,
        "DATA": r.data,
        "VALOR_SISPAG": float(r.valor_sispag),
        "STATUS": r.status,
        "TOTAL_ENCONTRADO": float(r.total_encontrado),
        "DIFERENCA": float(r.diferenca),
        "QTD_PAGAMENTOS": qtd,
        "TITULOS": titulos,
        "NOMES": nomes,
        "NF_NUM": nf_num,
        "VALORES": valores,
        "FORMA_PGTO": forma_pgto,
        "HISTORICOS": historicos,
    }


def gravar_auditoria(resultados: list[ResultadoSispag], caminho_saida: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Auditoria SISPAG"

    fonte_cabecalho = Font(name="Arial", size=10, bold=True)
    fonte_dado = Font(name="Arial", size=10)

    ws.append(COLUNAS_AUDITORIA)
    for celula in ws[1]:
        celula.font = fonte_cabecalho

    for r in resultados:
        linha = _linha_auditoria(r)
        ws.append([linha[col] for col in COLUNAS_AUDITORIA])

    for linha_celulas in ws.iter_rows(min_row=2):
        for celula in linha_celulas:
            celula.font = fonte_dado
        linha_celulas[1].number_format = "DD/MM/YYYY"  # DATA
        linha_celulas[2].number_format = "#,##0.00"  # VALOR_SISPAG
        linha_celulas[4].number_format = "#,##0.00"  # TOTAL_ENCONTRADO
        linha_celulas[5].number_format = "#,##0.00"  # DIFERENCA

    for coluna_celulas in ws.columns:
        largura = max((len(str(c.value)) for c in coluna_celulas if c.value is not None), default=0)
        ws.column_dimensions[coluna_celulas[0].column_letter].width = min(largura + 2, 60)

    caminho_saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho_saida)


def _imprimir_resumo(resultados: list[ResultadoSispag]) -> None:
    contagem = {STATUS_EXATO_1X1: 0, STATUS_EXATO_SOMA: 0, STATUS_AMBIGUO: 0, STATUS_NAO_ENCONTRADO: 0}
    for r in resultados:
        contagem[r.status] += 1

    print(f"Total de SISPAG FORNECEDORES processados: {len(resultados)}")
    for status, qtd in contagem.items():
        print(f"  {status}: {qtd}")

    ambiguos = [r for r in resultados if r.status == STATUS_AMBIGUO]
    if ambiguos:
        print("\nCasos AMBIGUO (SEQUENCIA, DATA, VALOR):")
        for r in ambiguos:
            print(f"  SEQ {r.sequencia} — {r.data} — R$ {r.valor_sispag}")

    nao_encontrados = [r for r in resultados if r.status == STATUS_NAO_ENCONTRADO]
    if nao_encontrados:
        print("\nCasos NAO_ENCONTRADO (SEQUENCIA, DATA, VALOR):")
        for r in nao_encontrados:
            print(f"  SEQ {r.sequencia} — {r.data} — R$ {r.valor_sispag}")


def executar(caminho_extrato: str, caminho_contas_pagas: str, caminho_saida: str) -> list[ResultadoSispag]:
    extrato = ler_extrato_padronizado(caminho_extrato)
    sispag_df = listar_sispag_fornecedores(extrato).sort_values("SEQUENCIA")

    sispags = [
        (
            int(linha["SEQUENCIA"]),
            linha["DATA"].date() if hasattr(linha["DATA"], "date") else linha["DATA"],
            valor_para_decimal(linha["VALOR"]),
        )
        for _, linha in sispag_df.iterrows()
    ]

    contas_pagas_df = ler_contas_pagas(caminho_contas_pagas)
    candidatos, _fora_do_escopo = montar_pagamentos(contas_pagas_df)

    resultados = conciliar(sispags, candidatos)

    gravar_auditoria(resultados, Path(caminho_saida))
    _imprimir_resumo(resultados)
    print(f"\nAuditoria gravada em: {caminho_saida}")

    return resultados


if __name__ == "__main__":
    if len(sys.argv) != 4:
        print(
            "Uso: python conciliar_sispag.py <extrato_padronizado.xlsx> "
            "<contas_pagas.xlsx> <auditoria_sispag.xlsx>"
        )
        sys.exit(1)
    try:
        executar(sys.argv[1], sys.argv[2], sys.argv[3])
    except (ExtratoPadronizadoError, ContasPagasError) as exc:
        print(f"ERRO: {exc}")
        sys.exit(1)
