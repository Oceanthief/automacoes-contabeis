"""Fase 1: converte o Extrato Itaú para o Excel padronizado de 8 colunas.

Escopo desta fase (ver Fase 0 para o desenho completo do projeto):
  - ler o extrato original;
  - localizar o cabeçalho dinamicamente;
  - transformar cada movimentação em uma linha DEBITO/CREDITO;
  - gravar um novo Excel — o arquivo original nunca é alterado.

Fora do escopo: qualquer leitura de Contas Pagas, e qualquer tentativa de
identificar SISPAG FORNECEDORES (isso é Fase 2).

Uso:
    python fase1_padronizar_extrato.py <extrato_origem.xlsx> <saida.xlsx>
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

from config import CONTA_PADRAO, COLUNAS_SAIDA
from normalizacao import data_para_datetime, valor_exportacao
from leitura.extrato_itau import ExtratoItauError, filtrar_movimentacoes, ler_extrato


class Fase1Error(Exception):
    """Erros de validação/negócio específicos desta fase."""


def transformar(movimentacoes: pd.DataFrame) -> pd.DataFrame:
    """Aplica a regra DEBITO/CREDITO linha a linha e monta o DataFrame de saída."""
    linhas = []
    for _, mov in movimentacoes.iterrows():
        valor_bruto = mov["Valor (R$)"]
        eh_pagamento = valor_bruto < 0

        linhas.append(
            {
                "col_sem_info_1": "",
                "debito": "" if eh_pagamento else CONTA_PADRAO,
                "credito": CONTA_PADRAO if eh_pagamento else "",
                "data": data_para_datetime(mov["Data"]),
                "valor": float(valor_exportacao(valor_bruto)),
                "col_sem_info_2": "",
                "descricao": mov["Lançamento"],
            }
        )

    saida = pd.DataFrame(linhas)
    saida["sequencia"] = range(1, len(saida) + 1)
    return saida


def validar(saida: pd.DataFrame) -> None:
    """Validações básicas antes de gravar — falha alto e claro, não silenciosamente."""
    if len(saida) == 0:
        raise Fase1Error("Nenhuma movimentação encontrada para exportar.")

    campos_obrigatorios = ["data", "valor", "descricao"]
    if saida[campos_obrigatorios].isna().any().any():
        linhas_ruins = saida[saida[campos_obrigatorios].isna().any(axis=1)]["sequencia"].tolist()
        raise Fase1Error(f"Linhas com DATA, VALOR ou DESCRICAO vazios: {linhas_ruins}")

    tem_debito = saida["debito"] != ""
    tem_credito = saida["credito"] != ""

    ambos = tem_debito & tem_credito
    if ambos.any():
        raise Fase1Error(
            f"DEBITO e CREDITO preenchidos simultaneamente na(s) linha(s): "
            f"{saida.loc[ambos, 'sequencia'].tolist()}"
        )

    nenhum = ~tem_debito & ~tem_credito
    if nenhum.any():
        raise Fase1Error(
            f"Linha sem DEBITO nem CREDITO preenchido: {saida.loc[nenhum, 'sequencia'].tolist()}"
        )

    sequencia_esperada = list(range(1, len(saida) + 1))
    if list(saida["sequencia"]) != sequencia_esperada:
        raise Fase1Error("SEQUENCIA não está em ordem contínua começando em 1.")


def gravar_excel(saida: pd.DataFrame, caminho_saida: Path) -> None:
    """Grava o Excel final com os dois cabeçalhos 'SEM INFORMACAO' literais."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Extrato Padronizado"

    fonte_cabecalho = Font(name="Arial", size=10, bold=True)
    fonte_dado = Font(name="Arial", size=10)

    rotulos = [rotulo for _, rotulo in COLUNAS_SAIDA]
    ws.append(rotulos)
    for celula in ws[1]:
        celula.font = fonte_cabecalho

    chaves = [chave for chave, _ in COLUNAS_SAIDA]
    for _, linha in saida.iterrows():
        ws.append([linha[chave] if chave != "sequencia" else int(linha[chave]) for chave in chaves])

    for linha_celulas in ws.iter_rows(min_row=2):
        for celula in linha_celulas:
            celula.font = fonte_dado
        linha_celulas[3].number_format = "DD/MM/YYYY"  # DATA
        linha_celulas[4].number_format = "#,##0.00"  # VALOR

    for coluna_celulas in ws.columns:
        largura = max((len(str(c.value)) for c in coluna_celulas if c.value is not None), default=0)
        ws.column_dimensions[coluna_celulas[0].column_letter].width = min(largura + 2, 40)

    caminho_saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho_saida)


def executar(caminho_origem: str, caminho_saida: str) -> pd.DataFrame:
    origem = Path(caminho_origem)
    destino = Path(caminho_saida)

    if destino.resolve() == origem.resolve():
        raise Fase1Error("O arquivo de saída não pode ser igual ao arquivo de origem.")

    df = ler_extrato(origem)
    movimentacoes = filtrar_movimentacoes(df)
    saida = transformar(movimentacoes)
    validar(saida)
    gravar_excel(saida, destino)

    print(f"OK: {len(saida)} movimentações exportadas para {destino}")
    return saida


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Uso: python fase1_padronizar_extrato.py <extrato_origem.xlsx> <saida.xlsx>")
        sys.exit(1)
    try:
        executar(sys.argv[1], sys.argv[2])
    except (ExtratoItauError, Fase1Error) as exc:
        print(f"ERRO: {exc}")
        sys.exit(1)
