"""Fase 6: validação contábil final, antes de qualquer geração de TXT.

Fluxo:
  1) lê saida_contabil.xlsx (Fase 5, intocado);
  2) aplica correcoes_manuais.xlsx por cima, numa cópia (Opção B combinada
     com o usuário) — nunca sobrescreve saida_contabil.xlsx;
  3) roda as 12 validações sobre o resultado;
  4) gera relatorio_validacao.xlsx com o veredito linha a linha e o
     STATUS_GERAL do arquivo inteiro.

Não gera TXT. Não corrige nada automaticamente.

Uso:
    python validar_saida_contabil.py <saida_contabil.xlsx> <correcoes_manuais.xlsx> \\
        <auditoria_granularizacao.xlsx> <relatorio_validacao.xlsx>
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

from leitura.correcoes_manuais import CorrecoesManuaisError, ler_correcoes_manuais
from validacao.aplicar_correcoes import AplicarCorrecoesError, aplicar_correcoes
from validacao.validador import (
    STATUS_APROVADO,
    STATUS_BLOQUEADO,
    STATUS_GERAL_APROVADO,
    calcular_totais,
    status_geral,
    validar_saida,
)


class ValidacaoError(Exception):
    pass


def _ler_saida_contabil(caminho: str) -> pd.DataFrame:
    caminho = Path(caminho)
    if not caminho.exists():
        raise ValidacaoError(f"Arquivo não encontrado: {caminho}")
    # DEBITO/CREDITO/TITULO são códigos/identificadores, não números — sem
    # dtype=str aqui o pandas derruba zero à esquerda (bug já visto 3x neste
    # projeto).
    return pd.read_excel(caminho, sheet_name=0, dtype={"DEBITO": str, "CREDITO": str, "TITULO": str})


def _gravar_relatorio(
    saida_validada: pd.DataFrame,
    contagem_categoria: dict,
    totais,
    status_geral_valor: str,
    caminho_saida: Path,
) -> None:
    wb = Workbook()
    fonte_cabecalho = Font(name="Arial", size=10, bold=True)
    fonte_dado = Font(name="Arial", size=10)

    def _escrever_aba(nome: str, df: pd.DataFrame, primeira: bool = False):
        ws = wb.active if primeira else wb.create_sheet(nome)
        ws.title = nome
        if df is None or len(df) == 0:
            ws.append(["(nenhum registro)"])
            return
        ws.append(list(df.columns))
        for c in ws[1]:
            c.font = fonte_cabecalho
        for _, linha in df.iterrows():
            ws.append([linha[col] for col in df.columns])
        for linha_celulas in ws.iter_rows(min_row=2):
            for celula in linha_celulas:
                celula.font = fonte_dado
        for coluna_celulas in ws.columns:
            largura = max((len(str(c.value)) for c in coluna_celulas if c.value is not None), default=0)
            ws.column_dimensions[coluna_celulas[0].column_letter].width = min(largura + 2, 60)

    # RESUMO
    resumo = pd.DataFrame(
        [
            ("STATUS_GERAL", status_geral_valor),
            ("Quantidade de movimentos bancários", totais.qtd_movimentos_bancarios),
            ("Quantidade de linhas contábeis", totais.qtd_linhas_contabeis),
            ("Linhas APROVADAS", int((saida_validada["STATUS_VALIDACAO"] == STATUS_APROVADO).sum())),
            ("Linhas BLOQUEADAS", int((saida_validada["STATUS_VALIDACAO"] == STATUS_BLOQUEADO).sum())),
            ("Total pagamentos (R$)", float(totais.total_pagamentos)),
            ("Total recebimentos (R$)", float(totais.total_recebimentos)),
            ("Total SISPAG granularizado (R$)", float(totais.total_sispag_granularizado)),
            ("Total bancário dos SISPAG (R$)", float(totais.total_bancario_sispag)),
            ("Total débitos (R$)", float(totais.total_debitos)),
            ("Total créditos (R$)", float(totais.total_creditos)),
        ],
        columns=["Métrica", "Valor"],
    )
    _escrever_aba("RESUMO", resumo, primeira=True)

    _escrever_aba("LINHAS_APROVADAS", saida_validada[saida_validada["STATUS_VALIDACAO"] == STATUS_APROVADO])
    _escrever_aba("ERROS_BLOQUEANTES", saida_validada[saida_validada["STATUS_VALIDACAO"] == STATUS_BLOQUEADO])
    _escrever_aba("SEM_DE_PARA", saida_validada[saida_validada["STATUS_DE_PARA"] == "SEM_DE_PARA"])
    _escrever_aba("AMBIGUOS", saida_validada[saida_validada["STATUS_DE_PARA"] == "AMBIGUO"])
    _escrever_aba(
        "SISPAG",
        saida_validada[saida_validada["DESCRICAO_ORIGINAL"] == "SISPAG FORNECEDORES"],
    )
    _escrever_aba(
        "DUPLICIDADES",
        saida_validada[saida_validada["MOTIVO_BLOQUEIO"].str.contains("PAGAMENTO_DUPLICADO", na=False)],
    )
    _escrever_aba(
        "VALIDACOES",
        pd.DataFrame(list(contagem_categoria.items()), columns=["Categoria", "Linhas com problema"]),
    )

    caminho_saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho_saida)


def _imprimir_resumo(saida_validada: pd.DataFrame, contagem_categoria: dict, totais, status_geral_valor: str) -> None:
    print(f"STATUS_GERAL: {status_geral_valor}")
    print(f"Linhas totais: {totais.qtd_linhas_contabeis}")
    print(f"Linhas APROVADAS: {(saida_validada['STATUS_VALIDACAO'] == STATUS_APROVADO).sum()}")
    print(f"Linhas BLOQUEADAS: {(saida_validada['STATUS_VALIDACAO'] == STATUS_BLOQUEADO).sum()}")
    print()
    print("Motivos de bloqueio por categoria:")
    for categoria, qtd in contagem_categoria.items():
        print(f"  {categoria}: {qtd}")
    print()
    print(f"Total pagamentos: R$ {totais.total_pagamentos}")
    print(f"Total recebimentos: R$ {totais.total_recebimentos}")
    print(f"Total SISPAG granularizado: R$ {totais.total_sispag_granularizado}")
    print(f"Total bancário dos SISPAG: R$ {totais.total_bancario_sispag}")

    bloqueadas = saida_validada[saida_validada["STATUS_VALIDACAO"] == STATUS_BLOQUEADO]
    if len(bloqueadas):
        print("\nLinhas bloqueadas (SEQUENCIA_CONTABIL — motivo):")
        for _, linha in bloqueadas.iterrows():
            print(f"  {linha['SEQUENCIA_CONTABIL']} — {linha['MOTIVO_BLOQUEIO']}")


def executar(
    caminho_saida_contabil: str,
    caminho_correcoes_manuais: str,
    caminho_auditoria_granularizacao: str,
    caminho_relatorio: str,
):
    saida_df = _ler_saida_contabil(caminho_saida_contabil)
    correcoes_df = ler_correcoes_manuais(caminho_correcoes_manuais)
    saida_corrigida = aplicar_correcoes(saida_df, correcoes_df)

    auditoria_granularizacao_df = pd.read_excel(caminho_auditoria_granularizacao, sheet_name=0)

    saida_validada, contagem_categoria = validar_saida(saida_corrigida, auditoria_granularizacao_df)
    totais = calcular_totais(saida_validada, auditoria_granularizacao_df)
    status = status_geral(saida_validada)

    _gravar_relatorio(saida_validada, contagem_categoria, totais, status, Path(caminho_relatorio))
    _imprimir_resumo(saida_validada, contagem_categoria, totais, status)
    print(f"\nOK: relatório gravado em {caminho_relatorio}")

    return saida_validada, status


if __name__ == "__main__":
    if len(sys.argv) != 5:
        print(
            "Uso: python validar_saida_contabil.py <saida_contabil.xlsx> "
            "<correcoes_manuais.xlsx> <auditoria_granularizacao.xlsx> <relatorio_validacao.xlsx>"
        )
        sys.exit(1)
    try:
        executar(*sys.argv[1:5])
    except (CorrecoesManuaisError, AplicarCorrecoesError, ValidacaoError) as exc:
        print(f"ERRO: {exc}")
        sys.exit(1)
