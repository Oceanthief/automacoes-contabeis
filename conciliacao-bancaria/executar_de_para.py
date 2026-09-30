"""Executa o motor de DE-PARA sobre saida_enriquecida_sispag.xlsx e gera
auditoria_de_para.xlsx.

Escopo: só auditoria — não escreve um novo extrato com as contas já
preenchidas (isso vira insumo da Fase 5, que monta a saída contábil final
combinando isto com a granularização dos SISPAG). SISPAG FORNECEDORES nunca
é classificado aqui: fica marcado AGUARDA_GRANULARIZACAO.

Uso:
    python executar_de_para.py <saida_enriquecida_sispag.xlsx> <de_para.xlsx> <auditoria_de_para.xlsx>
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from enriquecimento.classificacao_extrato import (
    PREFIXO_PGTO,
    PREFIXO_RECEBIMENTO,
    eh_pagamento,
    eh_recebimento,
)
from enriquecimento.motor_de_para import (
    STATUS_AMBIGUO,
    STATUS_JA_CLASSIFICADO,
    STATUS_MATCH_EXATO,
    STATUS_MATCH_REGRA,
    STATUS_SEM_DE_PARA,
    MotorDeParaError,
    carregar_regras,
    classificar,
)
from leitura.de_para import DeParaError, ler_de_para
from leitura.extrato_padronizado import ExtratoPadronizadoError, ler_extrato_padronizado
from normalizacao import normalizar_conta_contabil, normalizar_texto_classificacao

DESCRICAO_SISPAG_FORNECEDORES = "SISPAG FORNECEDORES"
STATUS_AGUARDA_GRANULARIZACAO = "AGUARDA_GRANULARIZACAO"

COLUNAS_AUDITORIA = [
    "SEQUENCIA",
    "DATA",
    "VALOR",
    "TIPO_MOVIMENTO",
    "DESCRICAO",
    "NOVA_DESCRICAO",
    "TEXTO_MATCH",
    "STATUS_DE_PARA",
    "CONTA_ENCONTRADA",
    "TIPO_MATCH_UTILIZADO",
    "CHAVE_MATCH_UTILIZADA",
    "PRIORIDADE_UTILIZADA",
    "HISTORICO_PADRAO_UTILIZADO",
    "OBSERVACAO",
]


def derivar_texto_match(nova_descricao: str) -> str:
    """Remove só os prefixos que o próprio sistema adicionou (Fase de
    enriquecimento), nunca palavras arbitrárias, e normaliza para comparação.
    """
    texto = "" if nova_descricao is None else str(nova_descricao)
    if texto.startswith(PREFIXO_PGTO):
        texto = texto[len(PREFIXO_PGTO):]
    elif texto.startswith(PREFIXO_RECEBIMENTO):
        texto = texto[len(PREFIXO_RECEBIMENTO):]
    return normalizar_texto_classificacao(texto)


def classificar_linha(linha, regras) -> dict:
    """Aplica a ordem de decisão desta fase a uma linha do extrato:
    1) já classificada por regra anterior (nunca sobrescreve);
    2) SISPAG FORNECEDORES (aguarda a próxima fase);
    3) motor de DE-PARA.
    """
    pagamento = eh_pagamento(linha["DEBITO"], linha["CREDITO"])
    recebimento = eh_recebimento(linha["DEBITO"], linha["CREDITO"])
    tipo_movimento = "PAGAMENTO" if pagamento else ("RECEBIMENTO" if recebimento else None)

    descricao = str(linha["DESCRICAO"]).strip()
    nova_descricao = "" if linha.get("NOVA DESCRICAO") is None else str(linha.get("NOVA DESCRICAO") or "")
    texto_match = derivar_texto_match(nova_descricao)

    conta_faltante_atual = normalizar_conta_contabil(
        linha["DEBITO"] if pagamento else linha["CREDITO"]
    )

    if conta_faltante_atual:
        return {
            "tipo_movimento": tipo_movimento or "",
            "texto_match": texto_match,
            "status": STATUS_JA_CLASSIFICADO,
            "conta_encontrada": conta_faltante_atual,
            "tipo_match_utilizado": "",
            "chave_match_utilizada": "",
            "prioridade_utilizada": "",
            "historico_padrao_utilizado": "",
            "observacao": "Conta já preenchida por regra anterior (ex.: tarifa bancária); DE-PARA não consultado.",
        }

    if descricao == DESCRICAO_SISPAG_FORNECEDORES:
        return {
            "tipo_movimento": tipo_movimento or "",
            "texto_match": texto_match,
            "status": STATUS_AGUARDA_GRANULARIZACAO,
            "conta_encontrada": "",
            "tipo_match_utilizado": "",
            "chave_match_utilizada": "",
            "prioridade_utilizada": "",
            "historico_padrao_utilizado": "",
            "observacao": "SISPAG conciliado — granularização por pagamento individual fica para a próxima fase.",
        }

    resultado = classificar(regras, tipo_movimento, texto_match)
    return {
        "tipo_movimento": tipo_movimento or "",
        "texto_match": texto_match,
        "status": resultado.status,
        "conta_encontrada": resultado.conta_encontrada or "",
        "tipo_match_utilizado": resultado.tipo_match_utilizado or "",
        "chave_match_utilizada": resultado.chave_match_utilizada or "",
        "prioridade_utilizada": resultado.prioridade_utilizada if resultado.prioridade_utilizada is not None else "",
        "historico_padrao_utilizado": resultado.historico_padrao_utilizado or "",
        "observacao": resultado.observacao,
    }


def processar(saida_df, regras) -> list[dict]:
    linhas = []
    for _, linha in saida_df.iterrows():
        resultado = classificar_linha(linha, regras)
        linhas.append(
            {
                "SEQUENCIA": int(linha["SEQUENCIA"]),
                "DATA": linha["DATA"],
                "VALOR": float(linha["VALOR"]),
                "TIPO_MOVIMENTO": resultado["tipo_movimento"],
                "DESCRICAO": linha["DESCRICAO"],
                "NOVA_DESCRICAO": linha.get("NOVA DESCRICAO", "") or "",
                "TEXTO_MATCH": resultado["texto_match"],
                "STATUS_DE_PARA": resultado["status"],
                "CONTA_ENCONTRADA": resultado["conta_encontrada"],
                "TIPO_MATCH_UTILIZADO": resultado["tipo_match_utilizado"],
                "CHAVE_MATCH_UTILIZADA": resultado["chave_match_utilizada"],
                "PRIORIDADE_UTILIZADA": resultado["prioridade_utilizada"],
                "HISTORICO_PADRAO_UTILIZADO": resultado["historico_padrao_utilizado"],
                "OBSERVACAO": resultado["observacao"],
            }
        )
    return linhas


def gravar_auditoria(linhas: list[dict], caminho_saida: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Auditoria DE-PARA"

    fonte_cabecalho = Font(name="Arial", size=10, bold=True)
    fonte_dado = Font(name="Arial", size=10)

    ws.append(COLUNAS_AUDITORIA)
    for celula in ws[1]:
        celula.font = fonte_cabecalho

    for l in linhas:
        ws.append([l[col] for col in COLUNAS_AUDITORIA])

    for linha_celulas in ws.iter_rows(min_row=2):
        for celula in linha_celulas:
            celula.font = fonte_dado
        linha_celulas[1].number_format = "DD/MM/YYYY"  # DATA
        linha_celulas[2].number_format = "#,##0.00"  # VALOR
        linha_celulas[8].number_format = "@"  # CONTA_ENCONTRADA como texto

    for coluna_celulas in ws.columns:
        largura = max((len(str(c.value)) for c in coluna_celulas if c.value is not None), default=0)
        ws.column_dimensions[coluna_celulas[0].column_letter].width = min(largura + 2, 70)

    caminho_saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho_saida)


def _imprimir_resumo(linhas: list[dict]) -> None:
    from collections import Counter

    contagem = Counter(l["STATUS_DE_PARA"] for l in linhas)

    print(f"Total de lançamentos analisados: {len(linhas)}")
    print(f"  JA_CLASSIFICADO:          {contagem.get(STATUS_JA_CLASSIFICADO, 0)}")
    print(f"  MATCH_EXATO:              {contagem.get(STATUS_MATCH_EXATO, 0)}")
    print(f"  MATCH_REGRA:              {contagem.get(STATUS_MATCH_REGRA, 0)}")
    print(f"  AMBIGUO:                  {contagem.get(STATUS_AMBIGUO, 0)}")
    print(f"  SEM_DE_PARA:              {contagem.get(STATUS_SEM_DE_PARA, 0)}")
    print(f"  AGUARDA_GRANULARIZACAO:   {contagem.get(STATUS_AGUARDA_GRANULARIZACAO, 0)}")

    ambiguos = [l for l in linhas if l["STATUS_DE_PARA"] == STATUS_AMBIGUO]
    if ambiguos:
        print("\nCasos AMBIGUO:")
        for l in ambiguos:
            print(f"  SEQ {l['SEQUENCIA']} — {l['TEXTO_MATCH']} — {l['OBSERVACAO']}")

    sem_de_para = [l for l in linhas if l["STATUS_DE_PARA"] == STATUS_SEM_DE_PARA]
    if sem_de_para:
        print("\nCasos SEM_DE_PARA:")
        for l in sem_de_para:
            print(f"  SEQ {l['SEQUENCIA']} — {l['TEXTO_MATCH']}")


def executar(caminho_saida_enriquecida: str, caminho_de_para: str, caminho_saida: str) -> list[dict]:
    saida_df = ler_extrato_padronizado(caminho_saida_enriquecida)
    de_para_df = ler_de_para(caminho_de_para)
    regras = carregar_regras(de_para_df)

    linhas = processar(saida_df, regras)
    gravar_auditoria(linhas, Path(caminho_saida))

    _imprimir_resumo(linhas)
    print(f"\nOK: auditoria gravada em {caminho_saida}")

    return linhas


if __name__ == "__main__":
    if len(sys.argv) != 4:
        print(
            "Uso: python executar_de_para.py <saida_enriquecida_sispag.xlsx> "
            "<de_para.xlsx> <auditoria_de_para.xlsx>"
        )
        sys.exit(1)
    try:
        executar(sys.argv[1], sys.argv[2], sys.argv[3])
    except (ExtratoPadronizadoError, DeParaError, MotorDeParaError) as exc:
        print(f"ERRO: {exc}")
        sys.exit(1)
