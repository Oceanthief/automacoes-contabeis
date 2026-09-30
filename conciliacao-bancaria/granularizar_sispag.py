"""Fase 5A: granulariza os SISPAG FORNECEDORES conciliados em linhas
contábeis individuais e monta a saída contábil unificada.

Reaproveita integralmente:
  - auditoria_de_para.xlsx já publicada pela Fase 4, para os lançamentos
    NORMAL (não-SISPAG) — contas não recalculadas aqui;
  - granularizacao.reconstrucao_sispag (que reinvoca o motor de conciliação
    já validado) para os pagamentos estruturados de cada SISPAG;
  - enriquecimento.motor_de_para para classificar cada pagamento individual.

Não reimplementa nenhuma dessas camadas. Não gera TXT.

Uso:
    python granularizar_sispag.py <extrato_padronizado.xlsx> <contas_pagas.xlsx> \\
        <auditoria_sispag.xlsx> <auditoria_de_para.xlsx> <de_para.xlsx> \\
        <saida_contabil.xlsx> <auditoria_granularizacao.xlsx>
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

from enriquecimento.motor_de_para import MotorDeParaError, carregar_regras
from granularizacao.granularizar import (
    STATUS_GRAN_ERRO,
    STATUS_GRAN_NAO_CONCILIADO,
    STATUS_GRAN_OK,
    TIPO_ORIGEM_NORMAL,
    DESCRICAO_SISPAG_FORNECEDORES,
    CONTA_BANCO,
    GranularizacaoError,
    LinhaContabil,
    granularizar_sispag,
)
from granularizacao.reconstrucao_sispag import (
    EstruturaSispagInsuficienteError,
    reconstruir_resultados_sispag,
)
from leitura.de_para import DeParaError, ler_de_para
from normalizacao import valor_para_decimal

COLUNAS_SAIDA_CONTABIL = [
    "SEQUENCIA_CONTABIL", "SEQUENCIA_BANCO", "DEBITO", "CREDITO", "DATA", "VALOR",
    "DESCRICAO_ORIGINAL", "NOVA_DESCRICAO", "TIPO_MOVIMENTO", "TIPO_ORIGEM",
    "STATUS_CONCILIACAO", "STATUS_DE_PARA", "TITULO", "FORNECEDOR", "NF_NUM",
    "HISTORICO", "FORMA_PGTO", "CONTA_ORIGEM",
]

COLUNAS_AUDITORIA_GRANULARIZACAO = [
    "SEQUENCIA_BANCO", "DATA", "VALOR_BANCO", "STATUS_CONCILIACAO",
    "QTD_PAGAMENTOS_ESPERADOS", "QTD_LINHAS_GERADAS", "TOTAL_GRANULARIZADO",
    "DIFERENCA", "STATUS_GRANULARIZACAO", "TITULOS", "FORNECEDORES", "OBSERVACAO",
]


def montar_linhas_normais(auditoria_de_para_df: pd.DataFrame) -> list[LinhaContabil]:
    """1 linha bancária NORMAL (não-SISPAG) -> 1 linha contábil, contas
    exatamente como a Fase 4 já decidiu."""
    linhas = []
    for _, linha in auditoria_de_para_df.iterrows():
        if str(linha["DESCRICAO"]).strip() == DESCRICAO_SISPAG_FORNECEDORES:
            continue  # SISPAG é tratado pela granularização

        tipo_movimento = str(linha["TIPO_MOVIMENTO"]).strip()
        conta_encontrada = linha.get("CONTA_ENCONTRADA", "")
        conta_encontrada = "" if pd.isna(conta_encontrada) else str(conta_encontrada)

        if tipo_movimento == "PAGAMENTO":
            debito, credito = conta_encontrada, CONTA_BANCO
        elif tipo_movimento == "RECEBIMENTO":
            debito, credito = CONTA_BANCO, conta_encontrada
        else:
            debito, credito = "", ""

        nova_descricao = linha.get("NOVA_DESCRICAO", "")
        nova_descricao = "" if pd.isna(nova_descricao) else str(nova_descricao)

        linhas.append(
            LinhaContabil(
                sequencia_contabil=0,
                sequencia_banco=int(linha["SEQUENCIA"]),
                debito=debito,
                credito=credito,
                data=linha["DATA"],
                valor=valor_para_decimal(linha["VALOR"]),
                descricao_original=str(linha["DESCRICAO"]),
                nova_descricao=nova_descricao,
                tipo_movimento=tipo_movimento,
                tipo_origem=TIPO_ORIGEM_NORMAL,
                status_conciliacao="",
                status_de_para=str(linha["STATUS_DE_PARA"]),
                titulo="",
                fornecedor="",
                nf_num="",
                historico="",
                forma_pgto="",
                conta_origem="",
            )
        )
    return linhas


def gravar_saida_contabil(linhas: list[LinhaContabil], caminho_saida: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Saida Contabil"

    fonte_cabecalho = Font(name="Arial", size=10, bold=True)
    fonte_dado = Font(name="Arial", size=10)

    ws.append(COLUNAS_SAIDA_CONTABIL)
    for celula in ws[1]:
        celula.font = fonte_cabecalho

    for l in linhas:
        ws.append(
            [
                l.sequencia_contabil, l.sequencia_banco, l.debito, l.credito, l.data, float(l.valor),
                l.descricao_original, l.nova_descricao, l.tipo_movimento, l.tipo_origem,
                l.status_conciliacao, l.status_de_para, l.titulo, l.fornecedor, l.nf_num,
                l.historico, l.forma_pgto, l.conta_origem,
            ]
        )

    for linha_celulas in ws.iter_rows(min_row=2):
        for celula in linha_celulas:
            celula.font = fonte_dado
        linha_celulas[2].number_format = "@"  # DEBITO como texto
        linha_celulas[3].number_format = "@"  # CREDITO como texto
        linha_celulas[4].number_format = "DD/MM/YYYY"  # DATA
        linha_celulas[5].number_format = "#,##0.00"  # VALOR

    for coluna_celulas in ws.columns:
        largura = max((len(str(c.value)) for c in coluna_celulas if c.value is not None), default=0)
        ws.column_dimensions[coluna_celulas[0].column_letter].width = min(largura + 2, 60)

    caminho_saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho_saida)


def gravar_auditoria_granularizacao(auditorias, caminho_saida: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Auditoria Granularizacao"

    fonte_cabecalho = Font(name="Arial", size=10, bold=True)
    fonte_dado = Font(name="Arial", size=10)

    ws.append(COLUNAS_AUDITORIA_GRANULARIZACAO)
    for celula in ws[1]:
        celula.font = fonte_cabecalho

    for a in auditorias:
        ws.append(
            [
                a.sequencia_banco, a.data, float(a.valor_banco), a.status_conciliacao,
                a.qtd_pagamentos_esperados, a.qtd_linhas_geradas, float(a.total_granularizado),
                float(a.diferenca), a.status_granularizacao, a.titulos, a.fornecedores, a.observacao,
            ]
        )

    for linha_celulas in ws.iter_rows(min_row=2):
        for celula in linha_celulas:
            celula.font = fonte_dado
        linha_celulas[1].number_format = "DD/MM/YYYY"
        linha_celulas[2].number_format = "#,##0.00"
        linha_celulas[6].number_format = "#,##0.00"
        linha_celulas[7].number_format = "#,##0.00"

    for coluna_celulas in ws.columns:
        largura = max((len(str(c.value)) for c in coluna_celulas if c.value is not None), default=0)
        ws.column_dimensions[coluna_celulas[0].column_letter].width = min(largura + 2, 70)

    caminho_saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho_saida)


def _imprimir_resumo(linhas_normais, linhas_sispag, auditorias, requer_class) -> None:
    from collections import Counter

    print(f"Movimentos bancários antes da granularização: {len(linhas_normais) + len(auditorias)}")
    print(f"Linhas contábeis depois da granularização: {len(linhas_normais) + len(linhas_sispag)}")
    print()

    qtd_1x1 = sum(1 for a in auditorias if a.status_conciliacao == "EXATO_1X1")
    qtd_soma = sum(1 for a in auditorias if a.status_conciliacao == "EXATO_SOMA")
    print(f"SISPAG EXATO_1X1: {qtd_1x1}")
    print(f"SISPAG EXATO_SOMA: {qtd_soma}")
    print(f"Total de pagamentos granularizados: {len(linhas_sispag)}")
    print()

    contagem_de_para = Counter(l.status_de_para for l in linhas_sispag)
    print(f"MATCH_EXATO: {contagem_de_para.get('MATCH_EXATO', 0)}")
    print(f"MATCH_REGRA: {contagem_de_para.get('MATCH_REGRA', 0)}")
    print(f"AMBIGUO: {contagem_de_para.get('AMBIGUO', 0)}")
    print(f"SEM_DE_PARA: {contagem_de_para.get('SEM_DE_PARA', 0)}")
    print()

    erros = [a for a in auditorias if a.status_granularizacao == STATUS_GRAN_ERRO]
    nao_conciliados = [a for a in auditorias if a.status_granularizacao == STATUS_GRAN_NAO_CONCILIADO]
    print(f"Erros de granularização (soma não fecha): {len(erros)}")
    for a in erros:
        print(f"  SEQUENCIA_BANCO {a.sequencia_banco}: diferença R$ {a.diferenca}")
    print(f"SISPAG não conciliados (sem pagamento para granularizar): {len(nao_conciliados)}")
    print()

    if requer_class:
        print(f"Casos REQUER_GRANULARIZACAO_CLASSIFICACAO: {len(requer_class)}")
        for c in requer_class:
            print(f"  {c}")
    else:
        print("Casos REQUER_GRANULARIZACAO_CLASSIFICACAO: 0")


def executar(
    caminho_extrato_padronizado: str,
    caminho_contas_pagas: str,
    caminho_auditoria_sispag: str,
    caminho_auditoria_de_para: str,
    caminho_de_para: str,
    caminho_saida_contabil: str,
    caminho_auditoria_granularizacao: str,
):
    resultados_sispag = reconstruir_resultados_sispag(
        caminho_extrato_padronizado, caminho_contas_pagas, caminho_auditoria_sispag
    )

    de_para_df = ler_de_para(caminho_de_para)
    regras_de_para = carregar_regras(de_para_df)

    linhas_sispag, auditorias, requer_class = granularizar_sispag(resultados_sispag, regras_de_para)

    auditoria_de_para_df = pd.read_excel(
        caminho_auditoria_de_para, sheet_name=0, dtype={"CONTA_ENCONTRADA": str}
    )
    linhas_normais = montar_linhas_normais(auditoria_de_para_df)

    todas_as_linhas = linhas_normais + linhas_sispag
    todas_as_linhas.sort(key=lambda l: (l.sequencia_banco, l.titulo))
    for i, linha in enumerate(todas_as_linhas, start=1):
        linha.sequencia_contabil = i

    gravar_saida_contabil(todas_as_linhas, Path(caminho_saida_contabil))
    gravar_auditoria_granularizacao(auditorias, Path(caminho_auditoria_granularizacao))

    _imprimir_resumo(linhas_normais, linhas_sispag, auditorias, requer_class)
    print(f"\nOK: saída contábil gravada em {caminho_saida_contabil}")
    print(f"OK: auditoria de granularização gravada em {caminho_auditoria_granularizacao}")

    return todas_as_linhas, auditorias, requer_class


if __name__ == "__main__":
    if len(sys.argv) != 8:
        print(
            "Uso: python granularizar_sispag.py <extrato_padronizado.xlsx> <contas_pagas.xlsx> "
            "<auditoria_sispag.xlsx> <auditoria_de_para.xlsx> <de_para.xlsx> "
            "<saida_contabil.xlsx> <auditoria_granularizacao.xlsx>"
        )
        sys.exit(1)
    try:
        executar(*sys.argv[1:8])
    except (EstruturaSispagInsuficienteError, GranularizacaoError, DeParaError, MotorDeParaError) as exc:
        print(f"ERRO: {exc}")
        sys.exit(1)
