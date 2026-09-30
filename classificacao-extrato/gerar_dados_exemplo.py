"""Gera um mês FICTÍCIO (agosto/2026) para rodar e testar a classificação.

Empresa, fornecedores, clientes, sócios e valores são inventados. O extrato
é gravado em .xls (o mesmo formato baixado do banco) e o cenário cobre todas
as regras do classificador:

  - transferências entre contas da empresa (filiais PR e SC, nos dois sentidos)
  - salários nas janelas do dia 6 e do dia 20 (±3 dias) e um FORA da janela
  - pró-labore vencendo a regra de salário quando os dois aparecem juntos
  - "Pago:" / "Recebido:" com DE-PARA, incluindo "o nome mais longo vence"
  - alias confirmado (grafia do banco diferente do cadastro)
  - cadastro com PREENCHER MANUAL e fornecedor sem cadastro -> pendências

Depois de gravar as entradas, o script repete a rotina real:
  preparar -> preencher a planilha de pendências -> finalizar

Uso:
    python gerar_dados_exemplo.py
"""
from __future__ import annotations

import shutil
import sys
import tempfile
from datetime import date
from pathlib import Path

import openpyxl
import xlwt

RAIZ = Path(__file__).resolve().parent
sys.path.insert(0, str(RAIZ / "src"))
sys.path.insert(0, str(RAIZ / "config"))

import configuracoes as cfg  # noqa: E402

DESTINO = RAIZ / "dados_exemplo"
ARQ_EXTRATO = "extrato_2026_08.xls"
ARQ_CONTROLE = "controle_fornecedores_clientes.xlsx"
ARQ_PENDENCIAS_PREENCHIDAS = "pendencias_preenchidas_2026_08.xlsx"

EMPRESA = cfg.NOME_EMPRESA

# (nome no Controle, conta débito, conta crédito, observação)
CONTROLE = [
    ("BANCO SANTANDER", "21850", cfg.CONTA_SANTANDER, "Juros e encargos bancários"),
    ("BANCO SANTANDER", "21850", cfg.CONTA_SANTANDER, "Cadastro duplicado (mesmas contas): não é ambiguidade"),
    ("BANCO SANTANDER - CONSORCIO", "13200", cfg.CONTA_SANTANDER, "Nome mais longo vence o genérico"),
    ("IOF", "21830", cfg.CONTA_SANTANDER, "Imposto sobre operações financeiras"),
    ("TARIFA", "21831", cfg.CONTA_SANTANDER, "Tarifas bancárias"),
    ("DISTRIBUIDORA ALFA LTDA", "21101", cfg.CONTA_SANTANDER, "Fornecedor de mercadorias"),
    ("TRANSPORTES BETA LTDA", "21102", cfg.CONTA_SANTANDER, "Frete"),
    ("J.P. DESPACHOS EXEMPLO", "21103", cfg.CONTA_SANTANDER, "Alias no config: 'J P DESPACHOS EXEMPLO'"),
    ("AP SERVICOS EXEMPLO", cfg.MARCADOR_PREENCHER_MANUAL, cfg.CONTA_SANTANDER, "Conta varia: sempre conferir"),
    ("CLIENTE GAMA COMERCIO LTDA", cfg.CONTA_SANTANDER, "11301", "Recebimento de cliente"),
    ("CLIENTE DELTA INDUSTRIA SA", cfg.CONTA_SANTANDER, "11302", "Recebimento de cliente"),
    ("PRO LABORE SOCIO A", cfg.CONTA_PRO_LABORE, cfg.CONTA_SANTANDER, "Pró-labore específico"),
    ("PRO LABORE", cfg.CONTA_PRO_LABORE, cfg.CONTA_SANTANDER, "Pró-labore genérico"),
]

# (dia, histórico, crédito, débito) — crédito = entrada, débito = saída
MOVIMENTOS = [
    (3, "IOF", None, 17.71),
    (3, "IOF", None, 4.36),
    (4, "Pago: 34330/08-48 - BANCO SANTANDER", None, 812.40),
    (5, f"Pago: 08001/08-12 - SALARIOS- {EMPRESA}", None, 1908.07),
    (5, f"Pago: 08002/08-12 - SALARIOS- {EMPRESA}", None, 1442.69),
    (5, "PRO LABORE SOCIO A", None, 2500.00),
    (5, f"TRANSFERENCIA DA CONTA SANTANDER {EMPRESA} - PR", 15000.00, None),
    (6, f"Pago: 08003/08-12 - SALARIOS- {EMPRESA} - PRO LABORE SOCIO B", None, 2500.00),
    (7, "Recebido: 55120/08 - CLIENTE GAMA COMERCIO LTDA", 4380.25, None),
    (7, "Pago: 10021/08-01 - DISTRIBUIDORA ALFA LTDA", None, 3120.90),
    (10, "Pago: 10022/08-01 - TRANSPORTES BETA LTDA", None, 640.00),
    (10, "TARIFA MENSALIDADE PACOTE SERVICOS", None, 89.90),
    (11, "Pago: 77001/08-02 - BANCO SANTANDER - CONSORCIO", None, 1250.00),
    (12, f"TRANSFERENCIA PARA CONTA SANTANDER {EMPRESA} - SC", None, 8000.00),
    (13, f"Pago: 08004/08-12 - SALARIOS- {EMPRESA}", None, 780.00),          # fora da janela -> pendência
    (14, "Recebido: 55121/08 - CLIENTE DELTA INDUSTRIA SA", 9870.10, None),
    (17, "Pago: 10023/08-01 - J P DESPACHOS EXEMPLO", None, 455.30),          # alias
    (18, "Pago: 10024/08-01 - AP SERVICOS EXEMPLO", None, 1200.00),           # PREENCHER MANUAL -> pendência
    (19, "Pago: 10025/08-01 - FORNECEDOR NOVO SEM CADASTRO LTDA", None, 318.75),  # sem cadastro -> pendência
    (20, f"Pago: 08005/08-12 - SALARIOS- {EMPRESA}", None, 1908.07),
    (21, f"TRANSFERENCIA PARA CONTA SANTANDER {EMPRESA} - PR", None, 5000.00),
    (24, "Recebido: 55122/08 - CLIENTE GAMA COMERCIO LTDA", 2210.00, None),
    (26, "Pago: 10026/08-01 - DISTRIBUIDORA ALFA LTDA", None, 1875.35),
    (28, "TARIFA SISPAG TITULO OUTRO BANCO", None, 12.50),
]

# Decisões que a contabilidade tomaria na planilha de pendências
# (chave = trecho do histórico)
DECISOES_MANUAIS = {
    "SALARIOS-": (cfg.CONTA_SALARIO_DIA_06, cfg.CONTA_SANTANDER, "Adiantamento pago fora da janela"),
    "AP SERVICOS EXEMPLO": ("21104", cfg.CONTA_SANTANDER, "Serviço de manutenção no mês"),
    "FORNECEDOR NOVO SEM CADASTRO": ("21105", cfg.CONTA_SANTANDER, "Novo fornecedor: incluir no Controle"),
}

SALDO_INICIAL = 1000.00


def gravar_extrato(destino: Path) -> None:
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Sheet1")
    formato_data = xlwt.easyxf(num_format_str="DD/MM/YYYY")
    formato_valor = xlwt.easyxf(num_format_str="#,##0.00")

    ws.write(0, 0, f"SANTANDER {EMPRESA} - AGOSTO 2026")
    for col, rotulo in enumerate(["Data", "Histórico", "Crédito", "Débito", "Saldo"]):
        ws.write(1, col, rotulo)

    saldo = SALDO_INICIAL
    ws.write(2, 0, date(2026, 7, 31), formato_data)
    ws.write(2, 1, "SALDO ANTERIOR")
    ws.write(2, 4, saldo, formato_valor)

    linha = 3
    for dia, historico, credito, debito in MOVIMENTOS:
        saldo = round(saldo + (credito or 0) - (debito or 0), 2)
        ws.write(linha, 0, date(2026, 8, dia), formato_data)
        ws.write(linha, 1, historico)
        if credito is not None:
            ws.write(linha, 2, credito, formato_valor)
        if debito is not None:
            ws.write(linha, 3, debito, formato_valor)
        ws.write(linha, 4, saldo, formato_valor)
        linha += 1

    ws.write(linha, 0, date(2026, 8, 31), formato_data)
    ws.write(linha, 1, "SALDO FINAL")
    ws.write(linha, 4, saldo, formato_valor)
    wb.save(str(destino))


def gravar_controle(destino: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Planilha1"
    ws.append(["FORNECEDOR/CLIENTE", "CONTA DEBITO", "CONTA CREDITO", None, None, "Obs"])
    for nome, deb, cred, obs in CONTROLE:
        ws.append([nome, deb, cred, None, None, obs])
    wb.save(destino)


def preencher_pendencias(caminho: Path) -> int:
    """Faz o papel da contabilidade: preenche as colunas manuais."""
    wb = openpyxl.load_workbook(caminho)
    ws = wb["Pendencias"]
    cabecalho = [c.value for c in ws[1]]
    idx = {nome: i + 1 for i, nome in enumerate(cabecalho)}
    col_hist = idx["HISTORICO_ORIGINAL"]
    preenchidas = 0
    for r in range(2, ws.max_row + 1):
        historico = str(ws.cell(r, col_hist).value or "")
        if not historico:
            continue
        decisao = next((d for chave, d in DECISOES_MANUAIS.items() if chave in historico.upper()), None)
        if decisao is None:
            raise SystemExit(f"Cenário inválido: pendência inesperada: {historico!r}")
        deb, cred, obs = decisao
        ws.cell(r, idx["CONTA_DEBITO_MANUAL"], deb)
        ws.cell(r, idx["CONTA_CREDITO_MANUAL"], cred)
        ws.cell(r, idx["OBSERVACAO_MANUAL"], obs)
        preenchidas += 1
    wb.save(caminho)
    return preenchidas


def main() -> None:
    from exportar_txt import gerar_txt_importacao
    from fluxo_mensal import processar_mes
    from lancamentos_finais import gerar_planilha_final
    from pendencias import nome_arquivo_pendencias

    if DESTINO.exists():
        shutil.rmtree(DESTINO)
    DESTINO.mkdir()
    extrato, controle = DESTINO / ARQ_EXTRATO, DESTINO / ARQ_CONTROLE
    gravar_extrato(extrato)
    gravar_controle(controle)

    with tempfile.TemporaryDirectory() as tmp:
        saida = Path(tmp)
        preparo = processar_mes(extrato, controle, 2026, 8, saida)
        if not preparo.bloqueado:
            raise SystemExit("Cenário inválido: o mês de exemplo deveria gerar pendências.")

        caminho_pendencias = saida / nome_arquivo_pendencias(2026, 8)
        qtd = preencher_pendencias(caminho_pendencias)
        shutil.copy(caminho_pendencias, DESTINO / ARQ_PENDENCIAS_PREENCHIDAS)

        final = gerar_planilha_final(extrato, controle, 2026, 8, saida)
        if not final.gerado:
            raise SystemExit(f"Cenário inválido: {final.motivo_bloqueio}")
        txt = gerar_txt_importacao(final.caminho_lancamentos, 2026, 8, saida)
        if not txt.gerado:
            raise SystemExit(f"Cenário inválido: {txt.motivo_falha}")

    total = round(sum(l.valor for l in final.linhas), 2)
    print(f"Dados fictícios gerados em {DESTINO.relative_to(RAIZ)}/")
    print(f"  movimentos: {len(preparo.classificacoes)} | status na preparação: {preparo.contagem_status}")
    print(f"  pendências preenchidas: {qtd}")
    print(f"  lançamentos finais: {len(final.linhas)} | total: {total:.2f}")


if __name__ == "__main__":
    main()
