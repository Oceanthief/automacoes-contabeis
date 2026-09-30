"""
Geração do TXT de importação para o ERP (Fase 6).

Conversão DETERMINÍSTICA Excel -> TXT. Lê exclusivamente a planilha final
já aprovada pela Fase 5 (`output/lancamentos_<empresa>_AAAA_MM.xlsx`) —
não volta a abrir o extrato bancário, não reclassifica nada, não
reprocessa nada.

`exemplo_importacao.txt` (auditado na Fase 1) foi usado só para
descobrir o CONTRATO TÉCNICO do layout — delimitador, aspas, encoding,
BOM, quebra de linha, casas decimais, ausência de cabeçalho. Nenhuma
conta ou histórico daquele arquivo é usada aqui; ele é de outro cliente.

Nunca deixa um .txt "meio certo" no output/: o arquivo só é movido para
o caminho final depois de passar pela validação de ida e volta (ler o
próprio TXT gerado de volta e comparar contra o Excel, linha a linha).
Se a validação falhar, nada fica em output/ além do que já existia.
"""

from __future__ import annotations

import csv
import sys
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Optional

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "config"))
import configuracoes as cfg

CODIGO_FIXO = cfg.CODIGO_ERP_FIXO
_COLUNAS_OFICIAIS_ESPERADAS = [
    "DATA", "CONTA DEBITO", "CONTA CREDITO", "NOVO HISTORICO", "VALOR", "SEQUENCIA DE LANCAMENTO",
]

# Contrato técnico do TXT, confirmado na Fase 1 a partir de
# exemplo_importacao.txt — replicado aqui, não copiado de lá. Fonte da
# verdade: config/configuracoes.py.
_ENCODING = cfg.TXT_ENCODING  # grava BOM automaticamente
_DELIMITADOR = cfg.TXT_DELIMITADOR
_QUOTECHAR = cfg.TXT_QUOTECHAR
_QUOTING = csv.QUOTE_ALL
_QUEBRA_DE_LINHA = cfg.TXT_QUEBRA_DE_LINHA
_FORMATO_DATA = cfg.TXT_FORMATO_DATA


class GeracaoTxtBloqueadaError(Exception):
    """Levantada quando o TXT não pode ser gerado — sempre com o motivo
    explícito, nunca silenciosamente."""


# ---------------------------------------------------------------------------
# Leitura da planilha final aprovada (Fase 5) — nunca o extrato bancário
# ---------------------------------------------------------------------------

@dataclass
class LancamentoAprovado:
    linha_excel: int  # número da linha em 'Lancamentos', para rastreio
    data: date
    conta_debito: str
    conta_credito: str
    novo_historico: str
    valor: float
    sequencia: int


def ler_lancamentos_aprovados(caminho_lancamentos: str | Path) -> list[LancamentoAprovado]:
    caminho = Path(caminho_lancamentos)
    if not caminho.exists():
        raise GeracaoTxtBloqueadaError(
            f"Planilha final aprovada não encontrada: {caminho}. Rode a Fase 5 primeiro — "
            "a Fase 6 nunca reprocessa o extrato bancário para chegar lá sozinha."
        )
    wb = openpyxl.load_workbook(caminho, data_only=True)
    if "Lancamentos" not in wb.sheetnames:
        raise GeracaoTxtBloqueadaError(f"Aba 'Lancamentos' não encontrada em {caminho}.")
    ws = wb["Lancamentos"]

    cabecalho = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    if cabecalho != _COLUNAS_OFICIAIS_ESPERADAS:
        raise GeracaoTxtBloqueadaError(
            f"Layout inesperado em {caminho}: cabeçalho {cabecalho} != {_COLUNAS_OFICIAIS_ESPERADAS}. "
            "A Fase 6 não tenta adivinhar colunas fora do padrão da Fase 5."
        )

    lancamentos: list[LancamentoAprovado] = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row[0] is None:
            continue
        data_bruta, conta_debito, conta_credito, historico, valor, sequencia = row

        if isinstance(data_bruta, datetime):
            data_final = data_bruta.date()
        elif isinstance(data_bruta, date):
            data_final = data_bruta
        else:
            raise GeracaoTxtBloqueadaError(f"Linha {i} de {caminho}: DATA não é uma data válida ({data_bruta!r}).")

        conta_debito = str(conta_debito).strip() if conta_debito is not None else ""
        conta_credito = str(conta_credito).strip() if conta_credito is not None else ""
        if not conta_debito or not conta_credito:
            raise GeracaoTxtBloqueadaError(f"Linha {i} de {caminho}: conta débito/crédito vazia.")

        historico = str(historico).strip() if historico is not None else ""
        if not historico:
            raise GeracaoTxtBloqueadaError(f"Linha {i} de {caminho}: NOVO HISTORICO vazio.")

        try:
            sequencia_int = int(sequencia)
        except (TypeError, ValueError):
            raise GeracaoTxtBloqueadaError(f"Linha {i} de {caminho}: SEQUENCIA DE LANCAMENTO inválida ({sequencia!r}).")

        lancamentos.append(LancamentoAprovado(
            linha_excel=i,
            data=data_final,
            conta_debito=conta_debito,
            conta_credito=conta_credito,
            novo_historico=historico,
            valor=round(float(valor), 2),
            sequencia=sequencia_int,
        ))

    if not lancamentos:
        raise GeracaoTxtBloqueadaError(f"Nenhuma linha de lançamento em {caminho}.")

    # ordem preservada exatamente como está no Excel (top a bottom) — a
    # SEQUENCIA de cada linha é só carregada adiante, nunca usada para
    # reordenar nem recalculada.
    return lancamentos


# ---------------------------------------------------------------------------
# Montagem e escrita do TXT
# ---------------------------------------------------------------------------

def montar_campos(l: LancamentoAprovado) -> list[str]:
    return [
        "",
        l.conta_debito,
        l.conta_credito,
        l.data.strftime(_FORMATO_DATA),
        f"{l.valor:.2f}",
        CODIGO_FIXO,
        l.novo_historico,
        str(l.sequencia),
    ]


def _escrever_txt(lancamentos: list[LancamentoAprovado], caminho: Path) -> None:
    with open(caminho, "w", encoding=_ENCODING, newline="") as f:
        writer = csv.writer(f, delimiter=_DELIMITADOR, quotechar=_QUOTECHAR,
                             quoting=_QUOTING, lineterminator=_QUEBRA_DE_LINHA)
        for l in lancamentos:
            writer.writerow(montar_campos(l))


# ---------------------------------------------------------------------------
# Validação obrigatória de ida e volta
# ---------------------------------------------------------------------------

@dataclass
class RelatorioValidacaoTxt:
    ok: bool
    checagens: list[tuple[str, bool, str]]


def validar_ida_e_volta(caminho_txt: Path, lancamentos_esperados: list[LancamentoAprovado]) -> RelatorioValidacaoTxt:
    """Relê o próprio TXT gerado e compara, campo a campo e linha a
    linha, contra o que a planilha final aprovada dizia. Nunca confia
    que a escrita funcionou só porque não deu exceção."""
    checagens: list[tuple[str, bool, str]] = []

    with open(caminho_txt, "rb") as f:
        bruto = f.read()
    tem_bom = bruto[:3] == b"\xef\xbb\xbf"
    checagens.append(("Arquivo gravado com BOM UTF-8 (mesmo padrão do arquivo de referência)", tem_bom, ""))
    tem_crlf = b"\r\n" in bruto
    checagens.append(("Quebra de linha é CRLF (mesmo padrão do arquivo de referência)", tem_crlf, ""))

    with open(caminho_txt, encoding=_ENCODING, newline="") as f:
        leitor = csv.reader(f, delimiter=_DELIMITADOR, quotechar=_QUOTECHAR)
        linhas_lidas = [linha for linha in leitor if linha]  # descarta eventual linha vazia final

    campos_invalidos = [i for i, linha in enumerate(linhas_lidas) if len(linha) != 8]
    checagens.append((
        "Todas as linhas têm exatamente 8 campos",
        len(campos_invalidos) == 0,
        "nenhuma" if not campos_invalidos else f"linhas com contagem errada: {campos_invalidos}",
    ))

    qtd_ok = len(linhas_lidas) == len(lancamentos_esperados)
    checagens.append((
        "Quantidade de registros no TXT = quantidade de lançamentos aprovados",
        qtd_ok,
        f"{len(linhas_lidas)} no TXT vs. {len(lancamentos_esperados)} aprovados",
    ))

    campo1_invalido: list[int] = []
    div_conta: list[tuple] = []
    div_data: list[tuple] = []
    div_valor: list[tuple] = []
    div_codigo: list[tuple] = []
    div_historico: list[tuple] = []
    div_sequencia: list[tuple] = []

    pode_comparar_linha_a_linha = qtd_ok and not campos_invalidos
    if pode_comparar_linha_a_linha:
        for idx, (linha_txt, esperado) in enumerate(zip(linhas_lidas, lancamentos_esperados)):
            vazio, debito, credito, data_txt, valor_txt, codigo, historico_txt, sequencia_txt = linha_txt

            if vazio != "":
                campo1_invalido.append(idx)
            if debito != esperado.conta_debito:
                div_conta.append((idx, "debito", debito, esperado.conta_debito))
            if credito != esperado.conta_credito:
                div_conta.append((idx, "credito", credito, esperado.conta_credito))

            try:
                data_lida = datetime.strptime(data_txt, _FORMATO_DATA).date()
            except ValueError:
                data_lida = None
            if data_lida != esperado.data:
                div_data.append((idx, data_txt, esperado.data.isoformat()))

            try:
                valor_lido = round(float(valor_txt), 2)
            except ValueError:
                valor_lido = None
            if valor_lido is None or abs(valor_lido - esperado.valor) > 0.005:
                div_valor.append((idx, valor_txt, esperado.valor))

            if codigo != CODIGO_FIXO:
                div_codigo.append((idx, codigo))

            if historico_txt != esperado.novo_historico:
                div_historico.append((idx, historico_txt, esperado.novo_historico))

            try:
                sequencia_lida = int(sequencia_txt)
            except ValueError:
                sequencia_lida = None
            if sequencia_lida != esperado.sequencia:
                div_sequencia.append((idx, sequencia_txt, esperado.sequencia))

    def _resumo(divergencias):
        return "nenhuma divergência" if not divergencias else str(divergencias[:5])

    checagens.append(("Primeiro campo vazio em todas as linhas", len(campo1_invalido) == 0,
                       "nenhuma" if not campo1_invalido else f"linhas: {campo1_invalido}"))
    checagens.append(("Contas (débito/crédito) batem linha a linha com o Excel", len(div_conta) == 0, _resumo(div_conta)))
    checagens.append(("Datas batem linha a linha com o Excel", len(div_data) == 0, _resumo(div_data)))
    checagens.append(("Valores batem linha a linha com o Excel (2 casas, ponto decimal)", len(div_valor) == 0, _resumo(div_valor)))
    checagens.append((f"Campo fixo '{CODIGO_FIXO}' presente em todas as linhas", len(div_codigo) == 0, _resumo(div_codigo)))
    checagens.append(("Histórico bate linha a linha com o Excel", len(div_historico) == 0, _resumo(div_historico)))
    checagens.append((
        "Sequência bate linha a linha com o Excel (a mesma validada na Fase 5, nunca recalculada aqui)",
        len(div_sequencia) == 0, _resumo(div_sequencia),
    ))

    ok = all(passou for _, passou, _ in checagens)
    return RelatorioValidacaoTxt(ok=ok, checagens=checagens)


# ---------------------------------------------------------------------------
# Ponto de entrada
# ---------------------------------------------------------------------------

@dataclass
class ResultadoTxt:
    gerado: bool
    caminho_txt: Optional[Path]
    relatorio_validacao: Optional[RelatorioValidacaoTxt]
    motivo_falha: Optional[str]
    quantidade_registros: int


def gerar_txt_importacao(
    caminho_lancamentos: str | Path,
    ano: int,
    mes: int,
    dir_saida: str | Path,
) -> ResultadoTxt:
    """Converte a planilha final aprovada (Fase 5) em TXT de importação.

    Escreve primeiro num arquivo temporário, valida ida e volta contra o
    Excel, e só então move para o caminho final
    `output/importacao_<empresa>_AAAA_MM.txt`. Se a validação falhar, apaga o temporário e não
    deixa nada em `output/` — a entrega é só o arquivo, nunca a
    importação em si no ERP (isso continua manual, fora desta fase).
    """
    dir_saida = Path(dir_saida)
    dir_saida.mkdir(parents=True, exist_ok=True)

    try:
        lancamentos = ler_lancamentos_aprovados(caminho_lancamentos)
    except GeracaoTxtBloqueadaError as e:
        return ResultadoTxt(False, None, None, str(e), 0)

    caminho_final = dir_saida / f"importacao_{cfg.PREFIXO_ARQUIVOS}_{ano:04d}_{mes:02d}.txt"
    caminho_temp = dir_saida / f".tmp_importacao_{cfg.PREFIXO_ARQUIVOS}_{ano:04d}_{mes:02d}.txt"

    _escrever_txt(lancamentos, caminho_temp)
    relatorio = validar_ida_e_volta(caminho_temp, lancamentos)

    if not relatorio.ok:
        caminho_temp.unlink(missing_ok=True)
        falhas = "; ".join(f"{n}: {d}" for n, passou, d in relatorio.checagens if not passou)
        motivo = f"Validação de ida e volta falhou — nenhum arquivo foi entregue. Falhas: {falhas}"
        return ResultadoTxt(False, None, relatorio, motivo, len(lancamentos))

    caminho_temp.replace(caminho_final)
    return ResultadoTxt(True, caminho_final, relatorio, None, len(lancamentos))
